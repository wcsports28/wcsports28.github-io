#!/usr/bin/env python3
"""
WCS Analytics — nightly NHL goalie updater.

1. Pulls tonight's starting goalies from Daily Faceoff.
2. Pulls MoneyPuck.com shot data and computes each starter's Sv% and
   GSAx/100 (goals saved above expected per 100 unblocked shots) vs
   LW, RW, C and D shooters.
3. Writes everything into the "Goalies" tab of WCS_NHL_Line_Stats.xlsx.
   The workbook's formulas then pick the weakest/strongest position and
   write the scouting note.

Usage:
    pip install requests pandas beautifulsoup4 openpyxl
    python wcs_goalie_update.py --workbook WCS_NHL_Line_Stats.xlsx
    python wcs_goalie_update.py --date 2026-10-08          # a specific slate

Data credit: shot data from MoneyPuck.com (credit required by MoneyPuck).
"""
import argparse, datetime as dt, io, re, sys, time, unicodedata, zipfile
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

SEASONS = [2025, 2026]  # MoneyPuck season label = start year (2025 = 2025-26)
MP_URL = "https://peter-tanner.com/moneypuck/downloads/shots_{season}.zip"
DFO_URL = "https://www.dailyfaceoff.com/starting-goalies/{date}"
UA = {"User-Agent": "Mozilla/5.0 (WCS Analytics goalie updater)"}
STATUSES = {"Confirmed", "Likely", "Expected", "Unconfirmed"}
POSITIONS = [("LW", "L"), ("RW", "R"), ("C", "C"), ("D", "D")]
NEED = {"goalieNameForShot", "playerPositionThatDidEvent", "xGoal", "goal",
        "shotWasOnGoal", "isPlayoffGame"}
FIRST_ROW, LAST_ROW = 6, 37  # Goalies tab input rows

TEAMS = {  # Daily Faceoff full name -> name used in the workbook
    "Anaheim Ducks": "Ducks", "Boston Bruins": "Bruins", "Buffalo Sabres": "Sabres",
    "Calgary Flames": "Flames", "Carolina Hurricanes": "Hurricanes",
    "Chicago Blackhawks": "Blackhawks", "Colorado Avalanche": "Avalanche",
    "Columbus Blue Jackets": "Blue Jackets", "Dallas Stars": "Stars",
    "Detroit Red Wings": "Red Wings", "Edmonton Oilers": "Oilers",
    "Florida Panthers": "Panthers", "Los Angeles Kings": "Kings", "Minnesota Wild": "Wild",
    "Montreal Canadiens": "Canadiens", "Montréal Canadiens": "Canadiens",
    "Nashville Predators": "Predators", "New Jersey Devils": "Devils",
    "New York Islanders": "Islanders", "New York Rangers": "Rangers",
    "Ottawa Senators": "Senators", "Philadelphia Flyers": "Flyers",
    "Pittsburgh Penguins": "Penguins", "San Jose Sharks": "Sharks",
    "Seattle Kraken": "Kraken", "St Louis Blues": "Blues", "St. Louis Blues": "Blues",
    "Tampa Bay Lightning": "Lightning", "Toronto Maple Leafs": "Maple Leafs",
    "Utah Mammoth": "Mammoth", "Vancouver Canucks": "Canucks",
    "Vegas Golden Knights": "Golden Knights", "Washington Capitals": "Capitals",
    "Winnipeg Jets": "Jets",
}


def norm(name: str) -> str:
    s = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", s.lower())


# ---------------------------------------------------------------- Daily Faceoff
def parse_starters(html: str) -> list[dict]:
    lines = [l.strip() for l in BeautifulSoup(html, "html.parser").get_text("\n").split("\n") if l.strip()]
    games, cur = [], None
    for i, line in enumerate(lines):
        m = re.fullmatch(r"(.+?) at (.+)", line)
        if m and m.group(1) in TEAMS and m.group(2) in TEAMS:
            cur = {"away": m.group(1), "home": m.group(2), "g": []}
            games.append(cur)
        elif cur is not None and line in STATUSES and len(cur["g"]) < 2:
            cur["g"].append((lines[i - 1], line))
    rows = []
    for g in games:
        if len(g["g"]) != 2:
            print(f"  ! skipped {g['away']} at {g['home']} (goalies not listed yet)")
            continue
        (ag, ast), (hg, hst) = g["g"]
        rows.append({"team": TEAMS[g["away"]], "goalie": ag, "status": ast, "opp": TEAMS[g["home"]]})
        rows.append({"team": TEAMS[g["home"]], "goalie": hg, "status": hst, "opp": TEAMS[g["away"]]})
    return rows


def get_starters(date: str, html_file: str | None) -> list[dict]:
    if html_file:
        html = Path(html_file).read_text(encoding="utf-8")
    else:
        r = requests.get(DFO_URL.format(date=date), headers=UA, timeout=30)
        r.raise_for_status()
        html = r.text
    rows = parse_starters(html)
    if not rows:
        sys.exit("ERROR: no starters parsed from Daily Faceoff — page layout may have changed.")
    return rows


# ---------------------------------------------------------------- MoneyPuck
def load_shots(cache: Path, offline: bool) -> pd.DataFrame:
    cache.mkdir(parents=True, exist_ok=True)
    frames = []
    for season in SEASONS:
        path = cache / f"shots_{season}.zip"
        stale = not path.exists() or time.time() - path.stat().st_mtime > 18 * 3600
        if stale and not offline:
            try:
                r = requests.get(MP_URL.format(season=season), headers=UA, timeout=180)
                r.raise_for_status()
                path.write_bytes(r.content)
                print(f"  downloaded MoneyPuck shots_{season}.zip")
            except requests.RequestException as e:
                print(f"  ! could not download season {season}: {e}")
        if not path.exists():
            continue
        with zipfile.ZipFile(path) as z:
            csv = next(n for n in z.namelist() if n.endswith(".csv"))
            df = pd.read_csv(z.open(csv), usecols=lambda c: c in NEED, low_memory=False)
        missing = NEED - set(df.columns)
        if missing:
            sys.exit(f"ERROR: MoneyPuck file is missing columns {missing} — check their data dictionary.")
        frames.append(df)
    if not frames:
        sys.exit("ERROR: no MoneyPuck shot data available.")
    df = pd.concat(frames, ignore_index=True)
    df = df[df["isPlayoffGame"] == 0].copy()
    df["gkey"] = df["goalieNameForShot"].map(norm)
    return df


def splits(shots: pd.DataFrame, goalie: str) -> dict | None:
    g = shots[shots["gkey"] == norm(goalie)]
    if g.empty:
        return None

    def calc(d):
        sog, goals, xg, n = d["shotWasOnGoal"].sum(), d["goal"].sum(), d["xGoal"].sum(), len(d)
        sv = round(1 - goals / sog, 3) if sog else None
        gsax = round((xg - goals) / n * 100, 2) if n else None
        return int(sog), sv, gsax

    out = {"total_sog": calc(g)[0], "overall": calc(g)[1:]}
    sogs = []
    for label, code in POSITIONS:
        sog, sv, gsax = calc(g[g["playerPositionThatDidEvent"] == code])
        out[label] = (sv, gsax)
        sogs.append(sog)
    out["min_sog"] = min(sogs)
    return out


# ---------------------------------------------------------------- Workbook
def write_workbook(path: Path, date: str, rows: list[dict], shots: pd.DataFrame):
    wb = load_workbook(path)
    ws = wb["Goalies"]
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for c in range(1, 17):
            ws.cell(r, c).value = None
    for i, row in enumerate(rows[: LAST_ROW - FIRST_ROW + 1]):
        r = FIRST_ROW + i
        ws.cell(r, 1, row["team"]); ws.cell(r, 2, row["goalie"])
        ws.cell(r, 3, row["status"]); ws.cell(r, 4, row["opp"])
        s = splits(shots, row["goalie"])
        if s is None:
            print(f"  ! no MoneyPuck history for {row['goalie']}")
            continue
        ws.cell(r, 5, s["total_sog"]); ws.cell(r, 6, s["min_sog"])
        col = 7
        for label, _ in POSITIONS:
            ws.cell(r, col, s[label][0]); ws.cell(r, col + 1, s[label][1]); col += 2
        ws.cell(r, 15, s["overall"][0]); ws.cell(r, 16, s["overall"][1])
        print(f"  {row['team']:<14} {row['goalie']:<22} {row['status']:<12} done")
    ws["B3"] = date
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates on open
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default="WCS_NHL_Line_Stats.xlsx")
    ap.add_argument("--date", default=dt.datetime.now(ZoneInfo("America/New_York")).strftime("%Y-%m-%d"))
    ap.add_argument("--cache", default=".mp_cache")
    ap.add_argument("--offline", action="store_true", help="use cached MoneyPuck files only")
    ap.add_argument("--html", help="parse a saved Daily Faceoff page instead of fetching (testing)")
    a = ap.parse_args()
    print(f"WCS goalie update for {a.date}")
    rows = get_starters(a.date, a.html)
    print(f"  {len(rows)} starters found")
    shots = load_shots(Path(a.cache), a.offline)
    write_workbook(Path(a.workbook), a.date, rows, shots)
    print("Done.")


if __name__ == "__main__":
    main()
